"""Operação de expansão: evidências persistentes e decisões versionadas."""
import base64
import hashlib
import io
import json
import uuid
from urllib.parse import urlparse
from flask import request, session, jsonify, send_file
from werkzeug.utils import secure_filename
import expansao_gestao as g

MARCOS = {'mobiliario':'Mobiliário','term_obra':'Término de obra','entrada_ti':'Entrada de TI','inauguracao':'Inauguração'}
MAX_ARQUIVO = 5 * 1024 * 1024

def migrar(db):
    with g.transacao(db) as (_,c):
        c.execute("CREATE TABLE IF NOT EXISTS expansao_operacao (filial TEXT PRIMARY KEY, dados TEXT NOT NULL DEFAULT '{}', versao INTEGER NOT NULL DEFAULT 0)")
        c.execute('CREATE TABLE IF NOT EXISTS expansao_evidencias (id TEXT PRIMARY KEY, filial TEXT NOT NULL, etapa TEXT NOT NULL, titulo TEXT NOT NULL, nome TEXT NOT NULL, conteudo TEXT NOT NULL, usuario TEXT NOT NULL, criado_em TEXT NOT NULL)')
        c.execute('CREATE INDEX IF NOT EXISTS idx_exp_evidencias ON expansao_evidencias(filial)')
        if db.IS_PG:
            c.execute("""CREATE OR REPLACE FUNCTION expansao_operacao_renomear() RETURNS trigger AS $$ BEGIN
            IF NEW.filial IS DISTINCT FROM OLD.filial THEN
              UPDATE expansao_operacao SET filial=NEW.filial WHERE filial=OLD.filial;
              UPDATE expansao_evidencias SET filial=NEW.filial WHERE filial=OLD.filial;
            END IF; RETURN NEW; END; $$ LANGUAGE plpgsql""")
            c.execute('DROP TRIGGER IF EXISTS expansao_operacao_rename ON acompanhamento_expansao')
            c.execute('CREATE TRIGGER expansao_operacao_rename AFTER UPDATE ON acompanhamento_expansao FOR EACH ROW EXECUTE FUNCTION expansao_operacao_renomear()')
        else:
            c.execute('''CREATE TRIGGER IF NOT EXISTS expansao_operacao_rename AFTER UPDATE ON acompanhamento_expansao WHEN NEW.filial!=OLD.filial BEGIN
            UPDATE expansao_operacao SET filial=NEW.filial WHERE filial=OLD.filial;
            UPDATE expansao_evidencias SET filial=NEW.filial WHERE filial=OLD.filial; END''')

def vazio():
    return dict(realizado={},remessas=[],aprovacao={},versao=0)

def carregar(db, filial, cur=None):
    if cur is None:
        with g.transacao(db) as (_,c): return carregar(db,filial,c)
    cur.execute(db.q('SELECT * FROM expansao_operacao WHERE filial=?'),(filial,))
    r=cur.fetchone(); d=vazio()
    if r: d.update(json.loads(r['dados'])); d['versao']=r['versao']
    cur.execute(db.q('SELECT id,etapa,titulo,nome,usuario,criado_em FROM expansao_evidencias WHERE filial=? ORDER BY criado_em,id'),(filial,))
    d['evidencias']=[dict(x) for x in cur.fetchall()]
    return d

def painel(db, linhas, gestao):
    """Carrega a operação de todas as filiais em duas consultas, sem anexos binários."""
    with g.transacao(db) as (_, c):
        c.execute('SELECT filial,dados,versao FROM expansao_operacao')
        registros={str(r['filial']):dict(r) for r in c.fetchall()}
        c.execute('SELECT id,filial,etapa,titulo,nome,usuario,criado_em FROM expansao_evidencias ORDER BY criado_em,id')
        evidencias={}
        for ev in c.fetchall():
            ev=dict(ev);filial=str(ev.pop('filial'));evidencias.setdefault(filial,[]).append(ev)
    ops={}
    for r in linhas:
        f=str(r['filial']);registro=registros.get(f);d=vazio()
        if registro:d.update(json.loads(registro['dados']));d['versao']=registro['versao']
        d['evidencias']=evidencias.get(f,[])
        d['situacao_aprovacao']=situacao_aprovacao(r,gestao[f]['plano'],d)
        # Listagens precisam apenas de marcos, remessas e situação. Detalhes no modal.
        ops[f]={k:d[k] for k in ('realizado','remessas','situacao_aprovacao')}
    return dict(operacoes=ops,hoje=g.hoje().isoformat())

def bloqueios(plano, evidencias):
    result=[]
    for t in plano.get('tarefas') or g.plano_vazio('')['tarefas']:
        if t['status'] not in ('Concluído','Não se aplica'): result.append(t['etapa']+': etapa pendente')
        elif t['status']=='Concluído' and not any(e['etapa']==t['etapa'] for e in evidencias): result.append(t['etapa']+': falta evidência')
        elif t['status']=='Não se aplica' and not t.get('observacao'): result.append(t['etapa']+': justifique Não se aplica')
    return result

def assinatura(r, plano, d):
    payload={'marcos':{k:r.get(k) for k in MARCOS},'checklist':plano.get('tarefas'), 'realizado':d['realizado'],'remessas':d['remessas'],'evidencias':[e['id'] for e in d['evidencias']]}
    return hashlib.sha256(json.dumps(payload,sort_keys=True,ensure_ascii=False).encode()).hexdigest()

def situacao_aprovacao(r,plano,d):
    a=d.get('aprovacao') or {}
    if a.get('status')=='Aprovada' and a.get('assinatura')!=assinatura(r,plano,d): return 'Revisão necessária'
    return a.get('status') or 'Pendente'

def validar(d):
    if not isinstance(d,dict): raise ValueError('Dados inválidos.')
    realizado=d.get('realizado',{}); remessas=d.get('remessas',[])
    if not isinstance(realizado,dict) or not isinstance(remessas,list) or len(remessas)>100: raise ValueError('Dados de operação inválidos.')
    def dt(v):
        if not v: return ''
        if not isinstance(v,str) or not g.data(v) or g.data(v).isoformat()!=v: raise ValueError('Data inválida.')
        return v
    real={k:dt(realizado.get(k)) for k in MARCOS}
    if any(g.data(v)>g.hoje() for v in real.values() if v): raise ValueError('Uma data realizada não pode estar no futuro.')
    envios=[]
    for r in remessas:
        if not isinstance(r,dict): raise ValueError('Remessa inválida.')
        item={k:str(r.get(k) or '').strip()[:lim] for k,lim in [('transportadora',120),('rastreio',120),('url',500),('observacao',1000),('recebido_por',120)]}
        item.update({k:dt(r.get(k)) for k in ('saida','previsao','entrega')})
        if not item['transportadora']: raise ValueError('Informe a transportadora de cada remessa.')
        if item['url'] and (urlparse(item['url']).scheme not in ('https','http') or not urlparse(item['url']).netloc): raise ValueError('Use um endereço HTTP ou HTTPS para rastreamento.')
        if item['entrega'] and (not item['saida'] or not item['recebido_por']): raise ValueError('Entrega exige saída e nome de quem recebeu.')
        if item['saida'] and item['entrega'] and item['entrega']<item['saida']: raise ValueError('A entrega não pode ser anterior à saída.')
        if any(g.data(item[k])>g.hoje() for k in ('saida','entrega') if item[k]): raise ValueError('Saída e entrega realizadas não podem estar no futuro.')
        item['id']=str(r.get('id') or uuid.uuid4().hex)[:64];envios.append(item)
    if len({r['id'] for r in envios})!=len(envios): raise ValueError('Remessas duplicadas.')
    return dict(realizado=real,remessas=envios)

def compras(linhas,gestao):
    result={}
    for r in linhas:
        if not g.aberto(r): continue
        for e in gestao[str(r['filial'])]['equipamentos']:
            if not e['falta']: continue
            key=(e['codigo'],e['descricao'])
            item=result.setdefault(key,dict(codigo=e['codigo'],descricao=e['descricao'],quantidade=0,filiais=[],necessidade=''))
            item['quantidade']+=e['falta'];item['filiais'].append(str(r['filial']))
            dt=g.data(r.get('entrada_ti')) or g.data(r.get('inauguracao'))
            if dt and (not item['necessidade'] or dt.isoformat()<item['necessidade']): item['necessidade']=dt.isoformat()
    return sorted(result.values(),key=lambda x:(x['necessidade'] or '9999',str(x['codigo'])))

def registrar(app,db,login_required,site_required,csrf_ok,estado):
    prefix='/api/acompanhamento-expansao'
    def acesso(registro_id, escrever=False):
        if escrever:
            if session.get('role') not in ('admin','gestor','operador'): return None,(jsonify(erro='Sem permissão para editar.'),403)
            if not csrf_ok(): return None,(jsonify(erro='Sessão expirada. Atualize a página.'),400)
        r=db.buscar_acompanhamento_expansao_por_id(registro_id)
        return (r,None) if r else (None,(jsonify(erro='Filial não encontrada.'),404))
    def plano(filial,c):
        c.execute(db.q('SELECT tarefas FROM expansao_planos WHERE filial=?'),(filial,));r=c.fetchone()
        return {'tarefas':json.loads(r['tarefas'])} if r else g.plano_vazio(filial)
    def travar(filial,c):
        # Mesmo lock do checklist para impedir aprovação concorrente de dados antigos.
        c.execute(db.q('INSERT INTO expansao_planos(filial) VALUES(?) ON CONFLICT(filial) DO NOTHING'),(filial,))
        c.execute(db.q('SELECT filial FROM expansao_planos WHERE filial=?'+(' FOR UPDATE' if db.IS_PG else '')),(filial,));c.fetchone()
        c.execute(db.q('INSERT INTO expansao_operacao(filial) VALUES(?) ON CONFLICT(filial) DO NOTHING'),(filial,))
    def gravar(filial,d,c,tipo,antes):
        persist={k:d[k] for k in ('realizado','remessas','aprovacao')}
        c.execute(db.q('UPDATE expansao_operacao SET dados=?,versao=versao+1 WHERE filial=?'),(json.dumps(persist,ensure_ascii=False),filial))
        g.evento(db,c,filial,tipo,antes,persist,session.get('username'))

    @app.route(prefix+'/operacao')
    @login_required
    @site_required('expansao')
    def expansao_operacao_painel():
        linhas,gestao=estado()
        saida=painel(db,linhas,gestao)
        saida.update(usuario=session.get('username'),compras=compras(g.filtrar(linhas,request.args),gestao))
        return jsonify(saida)

    @app.route(prefix+'/<int:registro_id>/operacao',methods=['GET','PUT'])
    @login_required
    @site_required('expansao')
    def expansao_operacao_filial(registro_id):
        r,err=acesso(registro_id,request.method=='PUT')
        if err:return err
        f=str(r['filial'])
        try:
            with g.transacao(db,request.method=='PUT') as (_,c):
                if request.method=='PUT': travar(f,c)
                d=carregar(db,f,c)
                if request.method=='PUT':
                    payload=request.get_json(silent=True); novo=validar(payload)
                    if payload.get('versao')!=d['versao']: return jsonify(erro='A operação foi alterada. Feche e reabra a filial antes de salvar.'),409
                    antes={k:d[k] for k in ('realizado','remessas','aprovacao')};d.update(novo);gravar(f,d,c,'operacao',antes);d['versao']+=1
                p=plano(f,c);d['revisao']=assinatura(r,p,d);d['bloqueios']=bloqueios(p,d['evidencias']);d['situacao_aprovacao']=situacao_aprovacao(r,p,d)
            return jsonify(d)
        except ValueError as e:return jsonify(erro=str(e)),400

    @app.route(prefix+'/<int:registro_id>/aprovacao',methods=['POST'])
    @login_required
    @site_required('expansao')
    def expansao_aprovar(registro_id):
        r,err=acesso(registro_id,True)
        if err:return err
        if session.get('role') not in ('admin','gestor'):return jsonify(erro='A aprovação exige perfil gestor ou administrador.'),403
        body=request.get_json(silent=True) or {}
        if not isinstance(body,dict) or body.get('status') not in ('Aprovada','Reprovada','Pendente'):return jsonify(erro='Decisão inválida.'),400
        motivo=str(body.get('motivo') or '').strip()[:1000]
        if not motivo:return jsonify(erro='Registre o parecer de TI.'),400
        f=str(r['filial'])
        with g.transacao(db,True) as (_,c):
            travar(f,c);d=carregar(db,f,c);p=plano(f,c)
            if body.get('versao')!=d['versao'] or body.get('revisao')!=assinatura(r,p,d):return jsonify(erro='Dados alterados. Reabra a filial para revisar.'),409
            impedimentos=bloqueios(p,d['evidencias'])
            if body['status']=='Aprovada' and impedimentos:return jsonify(erro='Não é possível aprovar: '+'; '.join(impedimentos)),409
            antes=d['aprovacao'];d['aprovacao']=dict(status=body['status'],motivo=motivo,usuario=session.get('username'),data=g.agora(),assinatura=assinatura(r,p,d))
            gravar(f,d,c,'aprovacao_ti',{'aprovacao':antes})
        return jsonify(ok=True)

    @app.route(prefix+'/<int:registro_id>/evidencias',methods=['POST'])
    @login_required
    @site_required('expansao')
    def expansao_evidencia_upload(registro_id):
        r,err=acesso(registro_id,True)
        if err:return err
        if request.content_length and request.content_length>MAX_ARQUIVO+65536:return jsonify(erro='O arquivo deve ter no máximo 5 MB.'),413
        etapa=request.form.get('etapa');titulo=request.form.get('titulo','').strip()[:200];arq=request.files.get('arquivo')
        if etapa not in g.ETAPAS or not titulo or not arq:return jsonify(erro='Informe etapa, título e arquivo.'),400
        conteudo=arq.read(MAX_ARQUIVO+1)
        if not conteudo or len(conteudo)>MAX_ARQUIVO:return jsonify(erro='O arquivo deve ter entre 1 byte e 5 MB.'),400
        nome=secure_filename(arq.filename or '')[:180] or 'evidencia'
        eid=uuid.uuid4().hex;f=str(r['filial'])
        with g.transacao(db,True) as (_,c):
            travar(f,c)
            c.execute(db.q('INSERT INTO expansao_evidencias VALUES(?,?,?,?,?,?,?,?)'),(eid,f,etapa,titulo,nome,base64.b64encode(conteudo).decode(),session.get('username') or 'sistema',g.agora()))
            c.execute(db.q('UPDATE expansao_operacao SET versao=versao+1 WHERE filial=?'),(f,))
            g.evento(db,c,f,'evidencia',{},dict(id=eid,etapa=etapa,titulo=titulo,nome=nome),session.get('username'))
        return jsonify(ok=True,id=eid),201

    @app.route(prefix+'/<int:registro_id>/evidencias/<eid>',methods=['GET','DELETE'])
    @login_required
    @site_required('expansao')
    def expansao_evidencia_arquivo(registro_id,eid):
        r,err=acesso(registro_id,request.method=='DELETE')
        if err:return err
        f=str(r['filial'])
        with g.transacao(db,request.method=='DELETE') as (_,c):
            if request.method=='DELETE':travar(f,c)
            c.execute(db.q('SELECT * FROM expansao_evidencias WHERE id=? AND filial=?'),(eid,f));ev=c.fetchone()
            if not ev:return jsonify(erro='Evidência não encontrada.'),404
            ev=dict(ev)
            if request.method=='DELETE':
                c.execute(db.q('DELETE FROM expansao_evidencias WHERE id=?'),(eid,));c.execute(db.q('UPDATE expansao_operacao SET versao=versao+1 WHERE filial=?'),(f,))
                g.evento(db,c,f,'exclusao_evidencia',{k:v for k,v in ev.items() if k!='conteudo'},{},session.get('username'))
                return jsonify(ok=True)
        response=send_file(io.BytesIO(base64.b64decode(ev['conteudo'])),mimetype='application/octet-stream',as_attachment=True,download_name=ev['nome'],max_age=0)
        response.headers['X-Content-Type-Options']='nosniff';response.headers['Cache-Control']='no-store';return response

    @app.route(prefix+'/compras.xlsx')
    @login_required
    @site_required('expansao')
    def expansao_compras_excel():
        from openpyxl import Workbook
        from openpyxl.styles import Font
        linhas,gestao=estado();rows=compras(g.filtrar(linhas,request.args),gestao)
        wb=Workbook();ws=wb.active;ws.title='Necessidade de compra'
        ws.append(['Código','Equipamento','Unidades faltantes','Filiais','Necessário até'])
        for x in rows:
            ws.append([x['codigo'],x['descricao'],x['quantidade'],', '.join(x['filiais']),x['necessidade']])
            for cell in ws[ws.max_row]:
                if isinstance(cell.value,str):cell.data_type='s'
        for cell in ws[1]:cell.font=Font(bold=True)
        for col,width in [('A',18),('B',45),('C',22),('D',40),('E',20)]:ws.column_dimensions[col].width=width
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        ws.append(['Estimativa após alocação global e reservas. Não cria pedido de compra.'])
        output=io.BytesIO();wb.save(output);output.seek(0)
        return send_file(output,as_attachment=True,download_name='expansao-necessidade-compra.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
