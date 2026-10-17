"""Testes funcionais isolados: nunca conecta ao Neon ou altera arquivos do usuário."""
import os
import sys
import tempfile
import unittest
import sqlite3
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor
from datetime import date

os.environ.pop('DATABASE_URL',None)
os.environ['REQUIRE_DATABASE_URL']='0'
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import db
TEMP=tempfile.TemporaryDirectory()
db.SQLITE_PATH=str(Path(TEMP.name)/'expansao.db')
import app as sistema
import expansao_gestao as exp

class GestaoTest(unittest.TestCase):
    def setUp(self):
        db.set_site_area('expansao')
        with exp.transacao(db,True) as (_,c):
            for t in ('expansao_reservas','expansao_planos','acompanhamento_expansao','expansao_eventos','itens','kit_padrao_loja','kit_filial_personalizado'):
                c.execute('DELETE FROM '+t)
            for filial,ti,inaug in [('001','2026-10-20','2026-10-25'),('002','2026-11-01','2026-11-10')]:
                c.execute(db.q('INSERT INTO acompanhamento_expansao(filial,descricao_filial,status_filial,uf,projeto,entrada_ti,inauguracao,atualizado_por) VALUES(?,?,?,?,?,?,?,?)'),(filial,'Loja de teste '+filial,'PENDENTE','SP','EXPANSÃO',ti,inaug,'teste'))
            c.execute("INSERT INTO kit_padrao_loja(codigo,descricao,quantidade,site_area) VALUES('PC','Computador',2,'expansao')")
            for item_id in (1,2,3):
                c.execute("INSERT INTO itens(id,codigo,descricao,qtde,tipo_estoque,status,site_area) VALUES(?,'PC','Computador','1','Expansão','Disponível','expansao')",(item_id,))
        self.client=sistema.app.test_client()
        with self.client.session_transaction() as s:
            s.update(user_id=1,username='teste',role='admin',site_area='expansao',ambiente_acesso='ambos',ambiente_selecionado=True,_csrf_token='token')
        self.ids={r['filial']:r['id'] for r in db.listar_acompanhamento_expansao()}
        self.headers={'X-CSRF-Token':'token'}

    def test_global_allocation_and_filter(self):
        r=self.client.get('/api/acompanhamento-expansao/gestao?mes=2026-11')
        self.assertEqual(r.status_code,200)
        d=r.json; self.assertEqual([r['filial'] for r in d['linhas']],['002'])
        self.assertEqual(d['resumo']['total'],1)
        self.assertEqual(d['gestao']['001']['sugeridos'],[1,2])
        self.assertEqual(d['gestao']['002']['faltantes'],1)
        self.assertEqual(d['gestao']['002']['sugeridos'],[3])

    def test_reservation_exclusion_dispatch_and_consumption(self):
        exp.reservar(db,'001',[1,2],'teste')
        with self.assertRaises(sqlite3.DatabaseError):
            with exp.transacao(db) as (_,c): c.execute("UPDATE itens SET status='Enviado',qtde='0',filial_destino='002' WHERE id=1")
        with self.assertRaises(sqlite3.DatabaseError):
            with exp.transacao(db) as (_,c): c.execute('DELETE FROM itens WHERE id=1')
        with self.assertRaises(sqlite3.DatabaseError):
            with exp.transacao(db) as (_,c): c.execute("UPDATE itens SET qtde='7' WHERE id=1")
        with exp.transacao(db) as (_,c): c.execute("UPDATE itens SET status='Enviado',qtde='0',filial_destino='001',atualizado_por='teste' WHERE id=1")
        _,rs=exp.planos_reservas(db);self.assertEqual([r['item_id'] for r in rs],[2])
        self.assertTrue(any(r['tipo']=='reserva_consumida' for r in exp.historico(db,'001')))
        exp.liberar(db,'001','teste')
        with exp.transacao(db) as (_,c): c.execute('DELETE FROM itens WHERE id=2')

    def test_concurrent_reservation_only_one_winner(self):
        def claim(f):
            try: exp.reservar(db,f,[1],'teste'); return True
            except ValueError: return False
        with ThreadPoolExecutor(2) as pool: results=list(pool.map(claim,['001','002']))
        self.assertEqual(sum(results),1)
        self.assertEqual(len(exp.planos_reservas(db)[1]),1)

    def test_plan_version_and_history(self):
        payload=exp.plano_vazio('001');payload['responsavel']='Alexandre';payload['tarefas'][0].update(responsavel='Equipe de rede',prazo='2026-10-19')
        url=f"/api/acompanhamento-expansao/{self.ids['001']}/gestao"
        self.assertEqual(self.client.put(url,json=payload,headers=self.headers).status_code,200)
        self.assertEqual(self.client.put(url,json=payload,headers=self.headers).status_code,409)
        d=self.client.get(url).json
        self.assertEqual(d['plano']['versao'],1);self.assertEqual(d['plano']['responsavel'],'Alexandre')
        self.assertTrue(any(r['tipo']=='checklist' and r['usuario']=='teste' for r in d['historico']))

    def test_permissions_csrf_and_cd(self):
        url=f"/api/acompanhamento-expansao/{self.ids['001']}/gestao"
        self.assertEqual(self.client.put(url,json=exp.plano_vazio('001')).status_code,400)
        with self.client.session_transaction() as s:s['role']='consulta'
        self.assertEqual(self.client.put(url,json=exp.plano_vazio('001'),headers=self.headers).status_code,403)
        self.assertEqual(self.client.get(url).status_code,200)
        with self.client.session_transaction() as s:s.update(site_area='estoque_cd',ambiente_acesso='estoque_cd')
        self.assertEqual(self.client.get('/api/acompanhamento-expansao/gestao').status_code,403)
        self.assertEqual(self.client.get('/api/acompanhamento-expansao/resumo-executivo.pdf').status_code,403)

    def test_risk_dates_and_unknowns(self):
        linhas=db.listar_acompanhamento_expansao();kits={x['filial']:db.listar_kit_padrao_loja() for x in linhas}
        result=exp.prever(linhas,db.listar_itens(),kits,{},[],sistema._equipamento_combina_kit,date(2026,10,26))
        self.assertEqual(result['001']['situacao'],'Atrasada')
        self.assertEqual(result['002']['situacao'],'Em risco')
        self.assertIsNone(exp.data('31/02/2026'))
        self.assertEqual(exp.data('17/08/26'),date(2026,8,17))
        self.assertEqual(exp.filtrar(linhas,{'mes':'2026-11','uf':'RJ'}),[])

    def test_api_reserve_requires_reviewed_forecast(self):
        url=f"/api/acompanhamento-expansao/{self.ids['001']}/reservas"
        self.assertEqual(self.client.post(url,json={'ids':[3]},headers=self.headers).status_code,409)
        self.assertEqual(self.client.post(url,json={'ids':[1,2]},headers=self.headers).status_code,200)
        self.assertEqual(self.client.delete(url,headers=self.headers).status_code,200)

    def test_audit_update_and_import(self):
        with self.client.session_transaction() as s:s['username']='autor_real'
        with sistema.app.test_request_context():
            sistema.session['username']='autor_real'
            with exp.transacao(db) as (_,c):c.execute("UPDATE acompanhamento_expansao SET inauguracao='2026-12-01',atualizado_por='antigo' WHERE filial='001'")
        hist=exp.historico(db,'001');alteracao=next(x for x in hist if x['tipo']=='alteracao')
        self.assertEqual(alteracao['usuario'],'autor_real');self.assertEqual(alteracao['antes']['inauguracao'],'2026-10-25');self.assertEqual(alteracao['depois']['inauguracao'],'2026-12-01')

    def test_rejected_stock_delete_keeps_connection_usable(self):
        exp.reservar(db,'001',[1],'teste')
        response=self.client.delete('/api/itens/1')
        self.assertEqual(response.status_code,409)
        self.assertTrue(db.buscar_item_por_id(1))
        exp.liberar(db,'001','teste')
        self.assertTrue(db.atualizar_item(1,{'qtde':'2'}))
        self.assertTrue(db.excluir_item(1))

    def test_backup_includes_new_data(self):
        import io, json, zipfile
        exp.reservar(db,'001',[1],'teste')
        response=self.client.get('/backup')
        self.assertEqual(response.status_code,200)
        with zipfile.ZipFile(io.BytesIO(response.data)) as z:
            d=json.loads(z.read('gestao_expansao.json'))
        self.assertEqual(d['expansao_reservas'][0]['filial'],'001')
        self.assertTrue(d['expansao_eventos'])

    def test_cd_stock_does_not_enter_expansion_history(self):
        with exp.transacao(db) as (_,c):
            c.execute("INSERT INTO itens(id,codigo,descricao,qtde,tipo_estoque,status,site_area,filial_destino) VALUES(99,'PC','Micro','1','CD','Disponível','estoque_cd','009')")
            c.execute("UPDATE itens SET qtde='2' WHERE id=99")
        self.assertEqual(exp.historico(db,'009'),[])

    def test_html_and_export_filters(self):
        html=self.client.get('/acompanhamento-expansao');self.assertEqual(html.status_code,200)
        self.assertIn(b'exp-dialog',html.data)
        excel=self.client.get('/export-acompanhamento-expansao?mes=2026-11')
        self.assertEqual(excel.status_code,200)
        import io
        from openpyxl import load_workbook
        wb=load_workbook(io.BytesIO(excel.data));self.assertEqual(wb['Acompanhamento'].max_row,2);self.assertEqual(wb['Acompanhamento']['A2'].value,'002')

if __name__=='__main__': unittest.main()
