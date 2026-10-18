import io
import json
import unittest
import zipfile
from datetime import timedelta
import test_expansao_gestao as base
from test_expansao_gestao import db, sistema, exp
import expansao_operacao as op

class OperacaoTest(unittest.TestCase):
    setUp=base.GestaoTest.setUp
    def url(self,f='001',suffix='operacao'):
        return f'/api/acompanhamento-expansao/{self.ids[f]}/{suffix}'
    def upload(self,etapa='Rede',content=b'evidencia de teste'):
        return self.client.post(self.url(suffix='evidencias'),data={'etapa':etapa,'titulo':'Teste','arquivo':(io.BytesIO(content),'teste.txt')},headers=self.headers)
    def aprovar(self,status='Aprovada',**extra):
        d=self.client.get(self.url()).json
        return self.client.post(self.url(suffix='aprovacao'),json={'status':status,'motivo':'Revisado por TI','versao':d['versao'],'revisao':d['revisao'],**extra},headers=self.headers)
    def test_panel_bulk_queries_and_combined_response(self):
        from unittest.mock import patch
        original=db.get_cursor
        calls=[]
        class CountCursor:
            def __init__(self,c):self.cursor=c
            def execute(self,sql,*args):calls.append(sql);return self.cursor.execute(sql,*args)
            def __getattr__(self,name):return getattr(self.cursor,name)
        linhas,gestao=sistema._expansao_estado()
        with patch.object(db,'get_cursor',side_effect=lambda conn:CountCursor(original(conn))):
            result=op.painel(db,linhas,gestao)
        self.assertEqual(len(calls),2)
        self.assertEqual(set(result['operacoes']),{'001','002'})
        response=self.client.get('/api/acompanhamento-expansao/gestao')
        self.assertIn('operacoes',response.json)
        self.assertIn('hoje',response.json)
        self.assertNotIn('exp-cronograma',self.client.get('/acompanhamento-expansao').text)

    def test_version_and_real_dates(self):
        d=self.client.get(self.url()).json
        d['realizado']={'mobiliario':'2026-01-01'}
        self.assertEqual(self.client.put(self.url(),json=d,headers=self.headers).status_code,200)
        self.assertEqual(self.client.put(self.url(),json=d,headers=self.headers).status_code,409)
        d=self.client.get(self.url()).json;d['realizado']['entrada_ti']=(exp.hoje()+timedelta(days=1)).isoformat()
        self.assertEqual(self.client.put(self.url(),json=d,headers=self.headers).status_code,400)
    def test_evidence_download_scope_delete_and_backup(self):
        r=self.upload();self.assertEqual(r.status_code,201);eid=r.json['id']
        self.assertEqual(self.client.get(self.url(suffix='evidencias/'+eid)).data,b'evidencia de teste')
        self.assertEqual(self.client.get(self.url('002','evidencias/'+eid)).status_code,404)
        backup=self.client.get('/backup')
        with zipfile.ZipFile(io.BytesIO(backup.data)) as z:d=json.loads(z.read('gestao_expansao.json'))
        self.assertEqual(d['expansao_evidencias'][0]['id'],eid)
        self.assertTrue(d['expansao_evidencias'][0]['conteudo'])
        self.assertEqual(self.client.delete(self.url(suffix='evidencias/'+eid),headers=self.headers).status_code,200)
        self.assertEqual(self.client.get(self.url(suffix='evidencias/'+eid)).status_code,404)
    def test_approval_requirements_and_invalidation(self):
        self.assertEqual(self.aprovar().status_code,409)
        p=exp.plano_vazio('001')
        for t in p['tarefas']:t['status']='Concluído'
        exp.salvar_plano(db,'001',p,'teste')
        self.assertEqual(self.aprovar().status_code,409)
        for etapa in exp.ETAPAS:self.assertEqual(self.upload(etapa).status_code,201)
        self.assertEqual(self.aprovar().status_code,200)
        self.assertEqual(self.client.get(self.url()).json['situacao_aprovacao'],'Aprovada')
        p['versao']=1;p['tarefas'][0]['observacao']='Mudou escopo'
        exp.salvar_plano(db,'001',p,'teste')
        self.assertEqual(self.client.get(self.url()).json['situacao_aprovacao'],'Revisão necessária')
        self.assertEqual(self.aprovar().status_code,200)
        eid=self.client.get(self.url()).json['evidencias'][0]['id']
        self.client.delete(self.url(suffix='evidencias/'+eid),headers=self.headers)
        self.assertEqual(self.client.get(self.url()).json['situacao_aprovacao'],'Revisão necessária')
    def test_approval_cannot_approve_unseen_checklist(self):
        p=exp.plano_vazio('001')
        for t in p['tarefas']:t.update(status='Não se aplica',observacao='Fora do escopo')
        exp.salvar_plano(db,'001',p,'teste')
        old=self.client.get(self.url()).json
        p['versao']=1;p['tarefas'][0]['observacao']='Alterado';exp.salvar_plano(db,'001',p,'teste')
        self.assertEqual(self.aprovar(revisao=old['revisao']).status_code,409)
        self.assertEqual(self.aprovar().status_code,200)
    def test_permissions(self):
        self.assertEqual(self.client.put(self.url(),json={}).status_code,400)
        with self.client.session_transaction() as s:s['role']='operador'
        self.assertEqual(self.aprovar().status_code,403)
        with self.client.session_transaction() as s:s['role']='consulta'
        self.assertEqual(self.upload().status_code,403)
        self.assertEqual(self.client.put(self.url(),json={},headers=self.headers).status_code,403)
        self.assertEqual(self.client.get(self.url()).status_code,200)
        with self.client.session_transaction() as s:s.update(site_area='estoque_cd',ambiente_acesso='estoque_cd')
        self.assertEqual(self.client.get(self.url()).status_code,403)
        self.assertEqual(self.client.get('/api/acompanhamento-expansao/operacao').status_code,403)
    def test_shipment_validation(self):
        d=self.client.get(self.url()).json
        d['remessas']=[dict(transportadora='Teste',saida='2026-01-01',entrega='2025-12-31',recebido_por='Teste')]
        self.assertEqual(self.client.put(self.url(),json=d,headers=self.headers).status_code,400)
        d['remessas'][0].update(entrega='2026-01-02',url='javascript:alert(1)')
        self.assertEqual(self.client.put(self.url(),json=d,headers=self.headers).status_code,400)
        d['remessas'][0]['url']='https://transportadora.example/rastreio'
        self.assertEqual(self.client.put(self.url(),json=d,headers=self.headers).status_code,200)
        self.assertEqual(self.client.get(self.url()).json['remessas'][0]['recebido_por'],'Teste')
    def test_upload_limits_and_invalid_payload(self):
        self.assertEqual(self.upload(content=b'x'*(op.MAX_ARQUIVO+1)).status_code,400)
        self.assertEqual(self.upload(etapa='Desconhecida').status_code,400)
        self.assertEqual(self.client.put(self.url(),json=[],headers=self.headers).status_code,400)
    def test_purchase_filter_and_rename(self):
        r=self.client.get('/api/acompanhamento-expansao/compras.xlsx?mes=2026-11')
        from openpyxl import load_workbook
        ws=load_workbook(io.BytesIO(r.data)).active
        self.assertEqual(ws['C2'].value,1);self.assertEqual(ws['D2'].value,'002')
        self.upload()
        with exp.transacao(db) as (_,c):c.execute("UPDATE acompanhamento_expansao SET filial='003' WHERE filial='001'")
        d=op.carregar(db,'003');self.assertEqual(len(d['evidencias']),1)
        self.assertEqual(op.carregar(db,'001')['evidencias'],[])

if __name__=='__main__':unittest.main()
